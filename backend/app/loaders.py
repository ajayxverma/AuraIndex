"""Extract plain text from uploaded files and split it into overlapping chunks."""
import csv
import io
from pathlib import PurePath

SUPPORTED_TYPES = {"pdf", "docx", "txt", "md", "csv", "xlsx", "html", "htm"}


def file_extension(filename: str) -> str:
    return PurePath(filename).suffix.lower().lstrip(".")


def extract_text(data: bytes, ext: str) -> str:
    if ext == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if ext == "docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext in ("txt", "md"):
        return data.decode("utf-8", errors="replace")
    if ext == "csv":
        reader = csv.reader(io.StringIO(data.decode("utf-8", errors="replace")))
        return "\n".join(", ".join(row) for row in reader)
    if ext == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"# Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    lines.append(", ".join("" if v is None else str(v) for v in row))
        return "\n".join(lines)
    if ext in ("html", "htm"):
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(data, "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()
        return soup.get_text("\n", strip=True)
    raise ValueError(f"Unsupported file type: {ext}")


def chunk_text(text: str, size: int, overlap: int) -> list[str]:
    text = " ".join(text.split())
    if not text:
        return []
    chunks, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            # Prefer breaking on a sentence/word boundary.
            cut = max(text.rfind(". ", start, end), text.rfind(" ", start, end))
            if cut > start + size // 2:
                end = cut + 1
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return [c for c in chunks if c]
